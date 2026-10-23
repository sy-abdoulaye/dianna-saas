"""Exports Excel (T-15, openpyxl) : une feuille, en-têtes figés, cellules neutralisées.

Comme pour le CSV (apps/core/csv_export.py), toute cellule qui commence par =, +, -, @, une
tabulation ou un retour chariot est préfixée d'une apostrophe, et toutes les cellules texte
sont enregistrées comme texte : openpyxl ne doit jamais en faire une formule.

Montants et heures : `Number`, enregistré comme nombre avec un format (« 1 234,56 € » ou
« 2,50 » dans un Excel français), pour que les sommes fonctionnent. Dans un CSV, la même
cellule s'écrit en texte français (« 42,05 € »)."""

import io
from collections.abc import Iterable
from dataclasses import dataclass
from decimal import Decimal
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter

from apps.core.csv_export import safe_cell
from apps.core.tokens import color

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class Styled:
    """Cellule colorée : `fill` et `ink` sont des noms de jetons (ex. "site-bg")."""

    text: str
    fill: str = ""
    ink: str = ""


# Codes de format Excel (indépendants de la langue : le séparateur de milliers et la
# virgule décimale sont ceux du poste qui ouvre le fichier).
MONEY_FORMAT = '#,##0.00\\ "€"'
DECIMAL_FORMAT = "#,##0.00"


@dataclass(frozen=True)
class Number:
    """Nombre d'un export : `money=True` pour un montant en euros."""

    value: Decimal
    money: bool = False

    def __str__(self) -> str:
        # Texte du CSV : virgule décimale, « € » après une espace insécable.
        if self.money:
            return f"{self.value:.2f}".replace(".", ",") + "\u00a0€"
        return str(self.value).replace(".", ",")


@lru_cache(maxsize=64)
def _fill(token: str) -> PatternFill:
    return PatternFill("solid", fgColor=color(token))


@lru_cache(maxsize=64)
def _font(token: str = "", bold: bool = False) -> Font:
    return Font(color=color(token) if token else None, bold=bold)


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _named_style(workbook, fill: str, ink: str) -> str:
    """Un style nommé par couple de jetons : une seule affectation par cellule (rapide en
    écriture seule sur un grand effectif)."""
    name = f"dianna-{fill or 'none'}-{ink or 'none'}"
    if name not in workbook.named_styles:
        style = NamedStyle(name=name)
        if fill:
            style.fill = _fill(fill)
        if ink:
            style.font = _font(ink)
        style.alignment = CENTER
        workbook.add_named_style(style)
    return name


def _cell(sheet, value, *, bold=False):
    """Cellule en écriture seule ; texte toujours neutralisé et enregistré comme texte."""
    if isinstance(value, Styled):
        cell = WriteOnlyCell(sheet, value=safe_cell(value.text))
        cell.data_type = "s"
        cell.style = _named_style(sheet.parent, value.fill, value.ink)
        return cell
    if isinstance(value, Number):
        cell = WriteOnlyCell(sheet, value=value.value)
        cell.number_format = MONEY_FORMAT if value.money else DECIMAL_FORMAT
        return cell
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return WriteOnlyCell(sheet, value=value)
    cell = WriteOnlyCell(sheet, value=safe_cell(value))
    cell.data_type = "s"
    if bold:
        cell.font = _font(bold=True)
    return cell


def xlsx_content(
    title: str, header: Iterable, rows: Iterable[Iterable], *, freeze: str = "A2", widths=None
) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(safe_cell(title)[:31] or "Export")
    sheet.freeze_panes = freeze
    for col, width in enumerate(widths or [], start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.append([_cell(sheet, value, bold=True) for value in header])
    for row in rows:
        sheet.append([_cell(sheet, value) for value in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
