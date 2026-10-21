"""Corrections de l'audit du lot 4 : coût des exports (L4-17), confidentialité du planning
(L4-16), dates extrêmes (L4-18), iCalendar coupé à la période."""

import csv
import io
from datetime import timedelta
from html import unescape
from pathlib import Path

import pytest
from django.conf import settings
from django.utils import timezone
from openpyxl import load_workbook

from apps.approvals.models import (
    LeaveRequest,
    OnCallDeclaration,
    RequestStatus,
    TeleworkRequest,
)
from apps.approvals.tests.conftest import next_weekday
from apps.core.models import AuditLog
from apps.core.ratelimit import EXPORT_LIMIT
from apps.core.tenancy import tenant_context
from apps.planning import exports
from apps.planning.tests.test_export import custom, download
from tests.factories import EmployeeFactory

pytestmark = pytest.mark.django_db
PDF_TOO_BIG = "Pour plus de personnes ou une période plus longue, utilisez l'export Excel."
LIMIT = "Vous avez fait 10 exports en moins de 10 minutes."


# ---------- a. Coût des exports ----------


def test_pdf_is_limited_to_one_month(client_for, company_a):
    monday = next_weekday()
    response = client_for(company_a.hr).get(
        f"/planning/export/?{custom(monday, monday + timedelta(days=31), fmt='pdf')}"
    )
    assert response.status_code == 422 and PDF_TOO_BIG in unescape(response.content.decode())
    assert (
        client_for(company_a.hr)
        .get(f"/planning/export/fichier/?{custom(monday, monday + timedelta(days=31), fmt='pdf')}")
        .status_code
        == 400
    )


@pytest.fixture
def crowd(company_a):
    """60 personnes de plus dans le service : plus d'une page du planning."""
    with tenant_context(company_a.company):
        for n in range(60):
            EmployeeFactory(
                company=company_a.company,
                service=company_a.service,
                site=company_a.site,
                last_name=f"Zz{n:02d}",
                user=None,
            )


def test_pdf_is_limited_to_one_page_of_the_planning(client_for, company_a, crowd, monkeypatch):
    monday = next_weekday()
    week = (monday, monday + timedelta(days=4))
    client = client_for(company_a.hr)
    refused = client.get(f"/planning/export/?{custom(*week, fmt='pdf')}")
    assert refused.status_code == 422 and PDF_TOO_BIG in unescape(refused.content.decode())
    captured = {}
    real = exports.pdf_context

    def spy(rows, first, last, company):
        captured["rows"] = rows
        return real(rows, first, last, company)

    monkeypatch.setattr(exports, "pdf_context", spy)
    download(client, custom(*week, fmt="pdf", page=2))
    names = [row.employee.last_name for row in captured["rows"]]
    assert len(names) == 15  # 65 personnes : la deuxième page en a 15
    assert all(name.startswith("Zz") for name in names)
    # Excel n'est pas limité.
    sheet = load_workbook(io.BytesIO(download(client, custom(*week, fmt="xlsx")).content)).active
    assert sheet.max_row == 1 + 65


def test_planning_export_button_keeps_the_page_shown(client_for, company_a, crowd):
    content = client_for(company_a.hr).get("/planning/?page=2").content.decode()
    assert 'href="/planning/export/?page=2"' in content


def test_exports_are_rate_limited_per_user(client_for, company_a):
    monday = next_weekday()
    client = client_for(company_a.hr)
    for _n in range(EXPORT_LIMIT):
        download(client, custom(monday, monday + timedelta(days=4)))
    refused = client.get(f"/planning/export/fichier/?{custom(monday, monday)}")
    assert refused.status_code == 429 and LIMIT in unescape(refused.content.decode())
    logged = AuditLog.all_companies.filter(company=company_a.company, action="planning.exported")
    assert logged.count() == EXPORT_LIMIT  # rien n'est généré au-delà
    preview = client.get(f"/planning/export/?{custom(monday, monday)}")
    assert LIMIT in unescape(preview.content.decode())
    # Le compteur est commun à tous les exports.
    hr = client.get("/teletravail/export/?mois=2026-11&format=csv")
    assert hr.status_code == 429
    campaign = company_a.requests["campaign"]
    summary = client.get(f"/conges/campagnes/{campaign.pk}/synthese/?format=csv")
    assert summary.status_code == 429
    modal = client.get("/teletravail/export/?mois=2026-11&format=csv", HTTP_HX_REQUEST="true")
    assert modal.status_code == 422 and LIMIT in unescape(modal.content.decode())
    # Une autre personne n'est pas touchée.
    other = client_for(company_a.admin)
    download(other, custom(monday, monday))


def test_gunicorn_has_an_explicit_timeout():
    dockerfile = (Path(settings.BASE_DIR) / "Dockerfile").read_text()
    assert '"--timeout", "30"' in dockerfile and '"--graceful-timeout", "30"' in dockerfile


# ---------- b. Confidentialité du planning (L4-16) ----------


@pytest.fixture
def nina(company_a):
    """Collègue du même service, dans l'équipe de Louise : congé d'été validé, télétravail en
    attente, télétravail validé, astreinte validée."""
    monday = next_weekday(offset_days=14)
    now = timezone.now()
    with tenant_context(company_a.company):
        person = EmployeeFactory(
            company=company_a.company,
            service=company_a.service,
            site=company_a.site,
            manager=company_a.manager,
            first_name="Nina",
            last_name="Petit",
        )
        from apps.approvals import workflow

        leave = LeaveRequest.objects.create(
            requester=person,
            leave_type=LeaveRequest.LeaveType.OTHER,
            start=monday,
            end=monday + timedelta(days=9),
            status=RequestStatus.APPROVED,
            submitted_at=now,
        )
        pending = TeleworkRequest.objects.create(
            requester=person, date=monday + timedelta(days=16), submitted_at=now
        )
        validated = TeleworkRequest.objects.create(
            requester=person,
            date=monday + timedelta(days=17),
            status=RequestStatus.APPROVED,
            submitted_at=now,
        )
        on_call = OnCallDeclaration.objects.create(
            requester=person,
            start=monday + timedelta(days=18),
            end=monday + timedelta(days=18),
            status=RequestStatus.APPROVED,
            submitted_at=now,
        )
        for request in (leave, pending, validated, on_call):
            workflow.sync_calendar(request)
    return person, monday


def nina_cells(client, monday, fmt="csv"):
    period = (monday, monday + timedelta(days=20))
    response = download(client, custom(*period, fmt=fmt))
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    return next(r for r in rows if r[0] == "Nina Petit")[3:]


def test_colleague_sees_absent_and_only_validated_requests(client_for, company_a, nina):
    _person, monday = nina
    cells = nina_cells(client_for(company_a.employee), monday)
    # Jours ouvrés : 0-4 et 7-11 (congé du lundi au mercredi suivant), puis 14-18…
    assert cells[0] == "Absent" and "Autre absence" not in cells
    assert cells[12] == "Sur site"  # télétravail en attente : invisible
    assert cells[13] == "Télétravail"
    assert cells[14] == "Sur site, astreinte"
    assert not any("en attente" in c for c in cells)


@pytest.mark.parametrize("who", ["manager", "hr", "admin"])
def test_manager_hr_and_admin_see_the_exact_type_and_pending(client_for, company_a, nina, who):
    _person, monday = nina
    cells = nina_cells(client_for(getattr(company_a, who)), monday)
    assert cells[0] == "Autre absence"
    assert cells[12] == "Télétravail (en attente)"


def test_everyone_sees_everything_about_themselves(client_for, company_a, nina):
    person, monday = nina
    from conftest import activate_totp  # noqa: F401 - Nina n'a pas de double authentification

    cells = nina_cells(client_for(person), monday)
    assert cells[0] == "Autre absence" and cells[12] == "Télétravail (en attente)"


def test_view_and_preview_follow_the_rule(client_for, company_a, nina):
    _person, monday = nina
    colleague = client_for(company_a.employee)
    view = unescape(colleague.get(f"/planning/?date={monday.isoformat()}").content.decode())
    assert "Absent" in view and "Autre absence" not in view
    month = unescape(
        colleague.get(f"/planning/?vue=mois&date={monday.isoformat()}").content.decode()
    )
    assert ">Ab<" in month and "Autre absence" not in month
    manager = unescape(
        client_for(company_a.manager).get(f"/planning/?date={monday.isoformat()}").content.decode()
    )
    assert "Autre absence" in manager
    preview = unescape(
        colleague.get(
            f"/planning/export/?{custom(monday, monday + timedelta(days=4))}"
        ).content.decode()
    )
    assert "Autre absence" not in preview


def test_excel_and_pdf_follow_the_rule(client_for, company_a, nina, monkeypatch):
    _person, monday = nina
    colleague = client_for(company_a.employee)
    period = (monday, monday + timedelta(days=20))
    sheet = load_workbook(
        io.BytesIO(download(colleague, custom(*period, fmt="xlsx")).content)
    ).active
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "Absent" in values and not any("Autre absence" in str(v) for v in values)
    assert not any("en attente" in str(v) for v in values if v)
    captured = {}
    real = exports.pdf_context

    def spy(rows, first, last, company):
        captured["labels"] = [c.label for r in rows for c in r.cells]
        return real(rows, first, last, company)

    monkeypatch.setattr(exports, "pdf_context", spy)
    download(colleague, custom(monday, monday + timedelta(days=4), fmt="pdf"))
    assert "Absent" in captured["labels"]
    assert not any("Autre absence" in label for label in captured["labels"])


def test_ics_follows_the_rule(client_for, company_a, nina):
    _person, monday = nina
    period = (monday, monday + timedelta(days=20))
    colleague = download(
        client_for(company_a.employee), custom(*period, fmt="ics")
    ).content.decode()
    assert "SUMMARY:Absent · Nina Petit" in colleague
    assert "Autre absence" not in colleague
    assert colleague.count("Nina Petit") == 2  # congé et télétravail validé, pas l'attente
    manager = download(client_for(company_a.manager), custom(*period, fmt="ics")).content.decode()
    assert "SUMMARY:Autre absence · Nina Petit" in manager
    assert manager.count("Nina Petit") == 3 and "STATUS:TENTATIVE" in manager


# ---------- c. Dates extrêmes (L4-18) ----------


@pytest.mark.parametrize("value", ["1500-01-01", "9999-12-31", "0001-01-01", "2101-01-01", "x"])
def test_planning_date_out_of_bounds_falls_back_to_today(client_for, company_a, value):
    from apps.planning import services

    today = timezone.localdate()
    for view, period in (("semaine", services.week_of), ("mois", services.month_of)):
        response = client_for(company_a.employee).get(f"/planning/?vue={view}&date={value}")
        assert response.status_code == 200
        expected = services.period_title(view, *period(today))
        assert expected in unescape(response.content.decode())


def test_planning_navigation_at_the_bounds_never_fails(client_for, company_a):
    client = client_for(company_a.employee)
    for url in ("/planning/?vue=mois&date=2100-12-01", "/planning/?vue=semaine&date=2000-01-01"):
        content = client.get(url).content.decode()
        assert client.get(url).status_code == 200
        for link in ("date=2101-01-01", "date=1999-12-20"):
            if link in content:
                assert client.get(f"/planning/?vue=mois&{link}").status_code == 200


@pytest.mark.parametrize(
    ("start", "end"),
    [("1999-12-31", "2000-01-02"), ("2100-12-30", "2101-01-02"), ("0001-01-01", "0001-01-02")],
)
def test_export_dates_out_of_bounds_are_form_errors(client_for, company_a, start, end):
    query = f"periode=personnalisee&debut={start}&fin={end}&format=csv"
    client = client_for(company_a.hr)
    response = client.get(f"/planning/export/?{query}")
    assert response.status_code == 422
    assert "Choisissez une date entre 2000 et 2100." in unescape(response.content.decode())
    assert client.get(f"/planning/export/fichier/?{query}").status_code == 400


def test_hr_month_export_out_of_bounds(client_for, company_a):
    response = client_for(company_a.hr).get(
        "/teletravail/export/?mois=1999-12&format=csv", HTTP_HX_REQUEST="true"
    )
    assert response.status_code == 422
    assert "Choisissez une date entre 2000 et 2100." in unescape(response.content.decode())


# ---------- d. iCalendar coupé à la période ----------


def test_ics_events_are_cut_to_the_requested_period(client_for, company_a, nina):
    _person, monday = nina  # congé du lundi au mercredi de la semaine suivante
    first, last = monday + timedelta(days=2), monday + timedelta(days=4)
    content = download(client_for(company_a.hr), custom(first, last, fmt="ics")).content.decode()
    assert f"DTSTART;VALUE=DATE:{first:%Y%m%d}" in content
    assert f"DTEND;VALUE=DATE:{last + timedelta(days=1):%Y%m%d}" in content
    assert f"{monday:%Y%m%d}" not in content
